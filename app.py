from flask import Flask, render_template, request, jsonify
import pandas as pd
import os
from openpyxl import load_workbook

app = Flask(__name__)

arquivo = "Produtos_modulacoes_colecao (2).xlsx"

df = pd.read_excel(arquivo)

df.columns = df.columns.str.strip().str.lower()

@app.route("/")
def home():
    return render_template("index.html")

# FEIRAS
@app.route("/feiras")
def feiras():
    feiras = sorted(df["colecao"].dropna().unique().tolist())
    return jsonify(feiras)

# PRODUTOS
@app.route("/produtos/<feira>")
def produtos(feira):
    produtos = df[df["colecao"] == feira]["desc_tecnica"].dropna().unique().tolist()
    produtos = sorted(produtos)
    return jsonify(produtos)

# MODULAÇÕES
@app.route("/modulacoes/<produto>")
def modulacoes(produto):
    modulacoes = df[df["desc_tecnica"] == produto]["caracteristica"].dropna().unique().tolist()
    modulacoes = sorted(modulacoes)
    return jsonify(modulacoes)

# VERSÕES
@app.route("/versoes/<modulacao>")
def versoes(modulacao):

    versoes = df[df["caracteristica"] == modulacao]["variavel"].dropna().unique().tolist()

    versoes = sorted(versoes)

    return jsonify(versoes)

# SALVAR
@app.route("/salvar", methods=["POST"])
def salvar():

    feira = request.form.get("feira")
    produto = request.form.get("produto")
    modulacao = request.form.get("modulacao")
    versao = request.form.get("versao")
    fase = request.form.get("fase")
    melhoria = request.form.get("melhoria")

    dados = {
        "Feira": feira,
        "Produto": produto,
        "Modulação": modulacao,
        "Versão": versao,
        "Fase": fase,
        "Tipo": melhoria
    }

    i = 1

    while True:

        descricao = request.form.get(f"descricao_{i}")

        if descricao is None:
            break

        dados[f"Descricao_{i}"] = descricao

        foto = request.files.get(f"foto_{i}")

        if foto and foto.filename != "":

            pasta = "static/uploads"

            os.makedirs(pasta, exist_ok=True)

            caminho = os.path.join(pasta, foto.filename)

            foto.save(caminho)

            dados[f"Foto_{i}"] = caminho

        else:
            dados[f"Foto_{i}"] = ""

        i += 1

    arquivo_excel = "registros.xlsx"

    novo_df = pd.DataFrame([dados])

    if os.path.exists(arquivo_excel):

        book = load_workbook(arquivo_excel)

        writer = pd.ExcelWriter(
            arquivo_excel,
            engine="openpyxl",
            mode="a",
            if_sheet_exists="overlay"
        )

        writer.book = book

        sheet = book.active

        startrow = sheet.max_row

        novo_df.to_excel(
            writer,
            index=False,
            header=False,
            startrow=startrow
        )

        writer.close()

    else:

        novo_df.to_excel(arquivo_excel, index=False)

    return "Salvo com sucesso!"

if __name__ == "__main__":
    app.run(host="0.0.0.0", port=10000)
